#!/usr/bin/env python3
"""
price_solve.py -- maximum purchase price that achieves each target return,
under both debt scenarios, using the SAME engine verify_model.py checks against
the workbook line by line (so the solve runs on verified mechanics, not a copy).

Targets
  (a) LP IRR = 8%            (LP clears the preferred return)
  (b) Levered IRR = 12%      (deal level)
  (c) Levered IRR = 15%      (deal level)

What moves with price (everything else held at the Assumptions-tab values):
  - Property tax: reassessed at price x 85% cost-of-sale factor x millage.
  - Closing costs: 1.5% of price.
  - Equity: the plug (Total Uses - debt).
  - Scenario B loan: LTV leg = 65% x price; DY/DSCR legs move only through NOI (tax).
  - Going-in (in-place) cap rate: Day-0 NOI at that price / price.
What does NOT move:
  - Scenario A assumed loan balances ($5.887M + $3.563M): an existing loan.
  - Rents, other opex, renovation capex, the assumption fee.
  - EXIT CAP (primary solve): frozen at the base-case market exit cap. The
    model's rule (in-place cap + 50 bps) would otherwise raise the exit cap
    whenever the buyer pays less, as if the market repriced because of what
    this buyer paid. The floating-exit-cap solve is reported as a secondary
    column so the difference is visible.
  - Refinance value (Scenario A): refi-year NOI / exit cap, so it moves only
    with NOI (tax), not with price directly.

Usage: python3 price_solve.py [workbook.xlsx]
"""
import sys
from scipy.optimize import brentq
import verify_model as V

PATH = sys.argv[1] if len(sys.argv) > 1 else "Parkview_Crossing_Acquisition_Model.xlsx"
I = V.load_inputs(PATH)
BASE = V.run_model(I)
EXIT_CAP_BASE = BASE["exit_cap"]
UNITS = I["units"]
LO, HI = 5_000_000, 25_000_000


def metric(price, scen, target, frozen):
    m = V.run_model(I, price=price, exit_cap=EXIT_CAP_BASE if frozen else None)
    S = m[scen]
    W = S["W"] if scen == "B" else m["W"]
    if S["equity"] <= 0:
        return None, m  # debt exceeds uses: not a coherent structure
    val = W["lp_irr"] if target == "lp" else S["irr"]
    return val, m


def floor_price(scen):
    """Lowest price at which equity is positive. Scenario A: the assumed loan is fixed, so
    equity = price x (1 + closing%) + reno capex + assumption fee - assumed loan = 0 solves in closed form."""
    if scen == "B":
        return LO
    A = BASE["A"]
    return (A["loan"] - BASE["reno_total"] - A["fee"]) / (1 + I["closing_pct"])


def solve(scen, target, rate, frozen):
    lo = max(LO, floor_price(scen) + 1_000)
    g = lambda p: metric(p, scen, target, frozen)[0] - rate
    try:
        glo, ghi = g(lo), g(HI)
    except TypeError:
        return None
    if glo * ghi > 0:
        return dict(unreachable=True, at_floor=glo + rate, floor=lo)
    p = brentq(g, lo, HI, xtol=1.0)
    val, m = metric(p, scen, target, frozen)
    S = m[scen]
    W = S["W"] if scen == "B" else m["W"]
    return dict(unreachable=False, price=p, ppu=p / UNITS, cap_inplace=m["cap_inplace"], cap_y1=m["cap_y1fwd"],
                deal_irr=S["irr"], lp_irr=W["lp_irr"], gp_irr=W["gp_irr"], equity=S["equity"],
                exit_cap=m["exit_cap"], disc=p / BASE["price"] - 1)


TARGETS = [("lp", 0.08, "(a) LP IRR = 8%"), ("deal", 0.12, "(b) Levered IRR = 12%"), ("deal", 0.15, "(c) Levered IRR = 15%")]
SCENS = [("A", "Scenario A — assumed loan (BASE CASE)"), ("B", "Scenario B — new debt (alternative)")]


def run():
    out = {}
    print(f"Base case at ${BASE['price']:,.0f} (${BASE['price']/UNITS:,.0f}/unit): "
          f"in-place cap {BASE['cap_inplace']:.2%}, exit cap {EXIT_CAP_BASE:.2%}, "
          f"A levered {BASE['A']['irr']:.2%} (LP {BASE['W']['lp_irr']:.2%}), "
          f"B levered {BASE['B']['irr']:.2%} (LP {BASE['B']['W']['lp_irr']:.2%})")
    print(f"Scenario A equity turns negative below ${floor_price('A'):,.0f} (assumed $9.45M loan > Total Uses).\n")
    for scen, slab in SCENS:
        print(slab)
        print(f"  {'Target':24s}{'Max price':>14s}{'$/unit':>10s}{'vs ask':>8s}{'In-place cap':>13s}"
              f"{'Deal IRR':>10s}{'LP IRR':>8s}{'GP IRR':>8s}{'| float-exit price':>20s}")
        for key, rate, tlab in TARGETS:
            r = solve(scen, key, rate, frozen=True)
            f = solve(scen, key, rate, frozen=False)
            out[(scen, tlab)] = (r, f)
            fstr = "unreachable" if (f is None or f["unreachable"]) else f"${f['price']:,.0f}"
            if r is None or r["unreachable"]:
                print(f"  {tlab:24s}{'UNREACHABLE above the structural floor':>63s}{fstr:>20s}")
                continue
            print(f"  {tlab:24s}{r['price']:>14,.0f}{r['ppu']:>10,.0f}{r['disc']:>8.1%}{r['cap_inplace']:>13.2%}"
                  f"{r['deal_irr']:>10.2%}{r['lp_irr']:>8.2%}{r['gp_irr']:>8.2%}{fstr:>20s}")
        print()
    return out


def check_summary(out, path):
    """Staleness guard: the Summary tab's Bid Price rows are typed values -- flag drift vs. a fresh solve."""
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["Summary"]
    key = {"A (base): LP IRR = 8%": ("A", "(a) LP IRR = 8%"), "A (base): Levered IRR = 12%": ("A", "(b) Levered IRR = 12%"),
           "A (base): Levered IRR = 15%": ("A", "(c) Levered IRR = 15%"), "B (alt): LP IRR = 8%": ("B", "(a) LP IRR = 8%"),
           "B (alt): Levered IRR = 12%": ("B", "(b) Levered IRR = 12%"), "B (alt): Levered IRR = 15%": ("B", "(c) Levered IRR = 15%")}
    stale = 0
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(row=r, column=1).value
        if lab in key:
            fresh = out[key[lab]][0]
            typed = ws.cell(row=r, column=2).value
            if fresh is None or fresh.get("unreachable") or abs(fresh["price"] - typed) > 1000:
                stale += 1
                print(f"  STALE Summary row '{lab}': typed {typed}, fresh {None if not fresh or fresh.get('unreachable') else round(fresh['price'])}")
    print("Summary Bid Price rows:", "STALE -- update build_model.py BID_ROWS" if stale else "match the fresh solve (within $1,000).")


if __name__ == "__main__":
    out = run()
    check_summary(out, PATH)
