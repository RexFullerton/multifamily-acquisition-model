#!/usr/bin/env python3
"""
verify_model.py -- independent recomputation of the Parkview Crossing model.

What it reads: ONLY literal input cells -- the blue inputs on the Assumptions
tab and on the Unit Mix tab (unit counts and rents live there per the model
spec). Every input is located by its row label, never by a row number, and
the script hard-fails if any cell it reads contains a formula. It never reads
a computed cell of the model to build its own numbers.

What it computes (its own engine, sharing no code with build_model.py):
10-year operating statement line by line (market GPR, loss-to-lease,
scheduled rent, vacancy, credit loss, concessions, other income, EGI, every
opex line, NOI, reserves, one-time capex, unlevered CF), the Day-0 NOI bridge
and cap rates, the tax-adjusted exit, Scenario B (new debt) sizing and
schedule, Scenario A (assumed first + supplemental loans, maturity refinance),
sources & uses for both, levered cash flows for both, the Scenario A
waterfall tier by tier, and IRR / equity multiple for everything.

What it then does: opens the recalculated workbook's cached values and
compares every line it computed, year by year, reporting any difference over
$1 (dollar lines) or 1 bp (rates, IRRs, multiples).

Usage:  python3 verify_model.py [recalculated.xlsx]
Exit code 0 = everything ties; 1 = at least one difference over tolerance.

The engine (load_inputs / run_model) is importable -- price_solve.py uses it.
"""
import sys
import openpyxl
import numpy_financial as npf

TOL_USD = 1.0
TOL_RATE = 0.0001  # 1 bp; also used for multiples (x)
NYEARS = 10        # operating model horizon
HOLD = 5           # structural hold (see Assumptions note)
UNIT_TYPES = ("Studio", "1BR/1BA", "2BR/1BA", "2BR/2BA")


# ---------------------------------------------------------------------------
# Input loading -- literal cells only
# ---------------------------------------------------------------------------
class InputIsFormula(Exception):
    pass


def _find_row(ws, label, col=1, exact=False, start=1, max_row=300):
    for r in range(start, max_row):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and (v.strip() == label if exact else label in v):
            return r
    raise KeyError(f"{ws.title}: no row label {'==' if exact else 'containing'} {label!r}")


def _literal(ws, r, c):
    v = ws.cell(row=r, column=c).value
    if isinstance(v, str) and v.startswith("="):
        raise InputIsFormula(f"{ws.title}!{ws.cell(row=r, column=c).coordinate} is a formula ({v}); "
                             f"verify_model.py only reads literal inputs")
    if not isinstance(v, (int, float)):
        raise ValueError(f"{ws.title}!{ws.cell(row=r, column=c).coordinate} is not numeric: {v!r}")
    return float(v)


def load_inputs(path):
    wb = openpyxl.load_workbook(path, data_only=False)  # formulas visible -> can assert literals
    a = wb["Assumptions"]

    def inp(label, exact=False):
        return _literal(a, _find_row(a, label, exact=exact), 3)

    I = {}
    I["price"] = inp("Purchase Price ($)")
    I["broker_cap"] = inp("Broker-Stated Cap Rate")
    I["closing_pct"] = inp("Closing Costs (% of Price)")
    I["millage"] = inp("Combined Millage Rate")
    I["seller_taxable"] = inp("Seller's Current TAXABLE Value")
    I["cos_factor"] = inp("Cost-of-Sale Factor")
    I["nonhs_cap"] = inp("Non-Homestead Annual Assessment Cap")
    I["reno_lines"] = [inp(x) for x in ("Kitchen (cabinet refacing", "Mini-split ductless AC",
                                        "Vinyl slider window", "Modern lighting fixtures", "LVP flooring",
                                        "Bathroom refresh", "Interior paint", "Turnover labor")]
    I["reno_contg_pct"] = inp("Contingency %", exact=True)
    I["reno_premium_mo"] = inp("Renovated-Unit Rent Premium over Market ($/month)")
    I["reno_y1_capture"] = inp("Year-1 Premium Capture %")
    I["vacancy"] = inp("Physical Vacancy %")
    I["credit_loss"] = inp("Credit Loss % (of GPR)")
    I["concessions"] = inp("Concessions % (of GPR)")
    I["other_income_mo"] = inp("Other Income ($/unit/month)")
    I["burnoff_pct"] = inp("Loss-to-Lease Burn-off % per Turnover")
    I["turnover_rate"] = inp("Annual Turnover Rate %")
    I["reno_headstart"] = inp("Prior-Owner-Renovated Units: Starting Capture")
    I["growth"] = [inp(f"Year {i} Market Rent Growth") for i in range(1, NYEARS + 1)]
    I["exp_growth"] = inp("General Expense Growth %")
    I["ins_growth"] = inp("Insurance Growth %")
    I["payroll"] = inp("Payroll (on-site")
    I["repairs"] = inp("Repairs & Maintenance", exact=True)
    I["turnover_cost"] = inp("Turnover / Make-Ready", exact=True)
    I["contract_svc"] = inp("Contract Services (landscaping")
    I["utilities"] = inp("Utilities (owner-paid")
    I["ins_base"] = inp("Insurance — Base Case")
    I["ins_roof"] = inp("Insurance — If Roof Replaced")
    I["mgmt_fee_pct"] = inp("Management Fee (% of EGI)")
    I["nav"] = inp("Non-Ad Valorem Assessments ($/unit")
    I["ga"] = inp("G&A / Admin", exact=True)
    I["marketing"] = inp("Marketing", exact=True)
    I["reserves"] = inp("Replacement Reserves (below NOI)")
    I["roof_toggle"] = inp("Roof Replacement Scenario Toggle")
    I["roof_cost"] = inp("Roof Replacement Cost")
    I["recert_year"] = inp("Recertification Year")
    I["recert_inspect"] = inp("Recertification Inspection Cost")
    I["recert_remed"] = inp("Recertification Remediation Contingency")
    I["sofr"] = inp("1-Month SOFR")
    I["spread"] = inp("Spread (bps over SOFR)")
    I["ltv"] = inp("Maximum LTV %")
    I["min_dscr"] = inp("Minimum DSCR")
    I["min_dy"] = inp("Minimum Debt Yield %")
    I["io_years"] = inp("Interest-Only Period (years)")
    I["amort_years"] = inp("Amortization (years)")
    I["a_first_bal"] = inp("First Mortgage Balance")
    I["a_first_rate"] = inp("First Mortgage Rate")
    I["a_mat"] = int(inp("First Mortgage Maturity"))
    I["a_supp_bal"] = inp("Supplemental Loan Balance")
    I["a_supp_rate"] = inp("Supplemental Loan Rate")
    I["a_io"] = inp("Both Loans Interest-Only")
    I["a_fee_pct"] = inp("Loan Assumption Fee")
    I["exit_spread_base"] = inp("Exit Cap Spread over Entry — Base Case")
    I["exit_spread_sens"] = inp("Exit Cap Spread over Entry — Sensitivity Ceiling")
    I["cos_exit"] = inp("Cost of Sale at Exit")
    I["pref"] = inp("Preferred Return %")
    I["t2_inv"] = inp("Tier 2 Investor Share")
    I["t2_promote"] = inp("Tier 2 GP Promote")
    I["t2_hurdle"] = inp("Tier 2 IRR Hurdle")
    I["t3_inv"] = inp("Tier 3 Investor Share")
    I["t3_promote"] = inp("Tier 3 GP Promote")
    I["gp_coinvest"] = inp("GP Co-Invest %")

    um = wb["Unit Mix"]
    types, r = [], 1
    while len(types) < len(UNIT_TYPES):  # stop before the Renovation Capex table, which reuses the names
        name = um.cell(row=r, column=1).value
        if name in UNIT_TYPES:
            types.append(dict(name=name, total=_literal(um, r, 2), classic=_literal(um, r, 3),
                              F=_literal(um, r, 6), H=_literal(um, r, 8)))
        r += 1
        if r > 50:
            raise KeyError("Unit Mix: did not find all 4 unit types")
    for t in types:  # recomputed here, NOT read from the model's column G / D
        t["renov"] = t["total"] - t["classic"]
        t["G"] = t["F"] + I["reno_headstart"] * (t["H"] - t["F"])
    I["types"] = types
    I["units"] = sum(t["total"] for t in types)
    I["classic_units"] = sum(t["classic"] for t in types)
    return I


# ---------------------------------------------------------------------------
# Engine
# ---------------------------------------------------------------------------
def _pmt(rate, n, pv):
    return pv * rate / (1 - (1 + rate) ** -n)


def _irr(cfs):
    try:
        v = npf.irr(cfs)
        return float(v)
    except Exception:
        return float("nan")


def waterfall(lev_cf, equity, I):
    """Pari passu waterfall, computed INDEPENDENTLY of the workbook's method.

    The workbook rolls two compounding hurdle balances forward year by year.
    Here each year's tier amounts are solved in closed form from the future
    value of the entire investor cash-flow history at each hurdle rate:
        shortfall_r(t) = max(0, E(1+r)^t - sum_{s<t} D_s (1+r)^(t-s))
    where D_s is everything investors (LP + GP co-invest) received in year s.
    Tier 1 = min(cash, pref shortfall); Tier 2 fills the Tier-2 shortfall at the
    investor share; Tier 3 takes the rest. Investor cash splits LP/GP pro rata
    to capital; the promote goes to GP. No state is carried except the history.
    Balance rows (for comparison only) are the same closed-form FVs restricted to
    the flows the workbook debits (Tier 1 for the pref balance; Tier 1 + investor
    Tier 2 for the Tier-2 balance)."""
    E, cash = equity, lev_cf[1:]
    p, h = I["pref"], I["t2_hurdle"]
    co = I["gp_coinvest"]
    D, T1, T2, T3, INV2, INV3, PR2, PR3 = [], [], [], [], [], [], [], []
    for t, av in enumerate(cash, start=1):
        fv_p = E * (1 + p) ** t - sum(D[s - 1] * (1 + p) ** (t - s) for s in range(1, t))
        t1 = min(av, max(0.0, fv_p))
        fv_h = E * (1 + h) ** t - sum(D[s - 1] * (1 + h) ** (t - s) for s in range(1, t)) - t1
        t2 = min(av - t1, max(0.0, fv_h) / I["t2_inv"])
        t3 = av - t1 - t2
        inv2, inv3 = t2 * I["t2_inv"], t3 * I["t3_inv"]
        T1.append(t1); T2.append(t2); T3.append(t3); INV2.append(inv2); INV3.append(inv3)
        PR2.append(t2 * I["t2_promote"]); PR3.append(t3 * I["t3_promote"])
        D.append(t1 + inv2 + inv3)
    n = len(cash)
    pref_bal = [E] + [E * (1 + p) ** t - sum(T1[s - 1] * (1 + p) ** (t - s) for s in range(1, t + 1)) for t in range(1, n + 1)]
    h_bal = [E] + [E * (1 + h) ** t - sum((T1[s - 1] + INV2[s - 1]) * (1 + h) ** (t - s) for s in range(1, t + 1))
                   for t in range(1, n + 1)]
    h_after_t1 = [h_bal[t - 1] * (1 + h) - T1[t - 1] for t in range(1, n + 1)]
    lp = [d * (1 - co) for d in D]
    gp_co = [d * co for d in D]
    promote = [PR2[i] + PR3[i] for i in range(n)]
    gp = [gp_co[i] + promote[i] for i in range(n)]
    lp_eq, gp_eq = E * (1 - co), E * co
    W = dict(avail=list(cash), t1=T1, t2=T2, t3=T3, inv2=INV2, inv3=INV3, pr2=PR2, pr3=PR3, inv_total=D,
             rem1=[cash[i] - T1[i] for i in range(n)], rem2=[cash[i] - T1[i] - T2[i] for i in range(n)],
             pref_bal=pref_bal, h_bal=h_bal, h_after_t1=h_after_t1, lp=lp, gp_co=gp_co, promote=promote, gp=gp,
             lp_eq=lp_eq, gp_eq=gp_eq, promote_total=sum(promote))
    W["lp_irr"], W["gp_irr"] = _irr([-lp_eq] + lp), _irr([-gp_eq] + gp)
    W["lp_em"], W["gp_em"] = sum(lp) / lp_eq, sum(gp) / gp_eq
    # structural assertions that hold for ANY correct pari passu waterfall
    for i in range(n):
        assert abs(lp[i] + gp[i] - cash[i]) < 1e-6, "LP + GP must equal available cash"
    if W["promote_total"] < 1e-9:
        assert abs(W["lp_irr"] - _irr(lev_cf)) < 1e-9 and abs(W["gp_irr"] - _irr(lev_cf)) < 1e-9, \
            "no promote => LP IRR = GP IRR = deal IRR"
    return W


def run_model(I, price=None, exit_cap=None, refi_basis="value"):
    """Full model. `price` overrides the purchase price (for bid-price solving):
    reassessed tax, closing costs, Scenario B loan sizing (LTV on price; DY/DSCR on
    NOI, which moves with tax) and equity move with it. The assumed Scenario A loan
    balances, rents, and all other opex do not.
    `exit_cap` freezes the exit cap (bid-price solve holds the market exit cap fixed
    rather than letting it float with the buyer's own entry cap). Default = the
    model rule: in-place cap at `price` + base spread."""
    P = I["price"] if price is None else price
    Y = range(NYEARS)
    ex = [(1 + I["exp_growth"]) ** y for y in Y]
    u = I["units"]
    rate = I["sofr"] + I["spread"]
    m = {}

    # rent tracks
    mk, cl, pr = {}, {}, {}
    for t in I["types"]:
        a, b, c = [], [], []
        for y in Y:
            g = I["growth"][y]
            mv = t["H"] * (1 + g) if y == 0 else a[-1] * (1 + g)
            cv = (t["F"] * (1 - I["reno_y1_capture"]) + (mv + I["reno_premium_mo"]) * I["reno_y1_capture"]
                  if y == 0 else b[-1] * (1 + g))
            base = t["G"] if y == 0 else c[-1]
            pv = base + I["turnover_rate"] * I["burnoff_pct"] * (mv - base)
            a.append(mv); b.append(cv); c.append(pv)
        mk[t["name"]], cl[t["name"]], pr[t["name"]] = a, b, c
    m["track_market"], m["track_classic"], m["track_prior"] = mk, cl, pr

    m["market_gpr"] = [sum(t["total"] * mk[t["name"]][y] for t in I["types"]) * 12 for y in Y]
    m["sched"] = [sum(t["classic"] * cl[t["name"]][y] + t["renov"] * pr[t["name"]][y] for t in I["types"]) * 12
                  for y in Y]
    m["ltl"] = [m["sched"][y] - m["market_gpr"][y] for y in Y]
    m["vac"] = [-x * I["vacancy"] for x in m["sched"]]
    m["cl"] = [-x * I["credit_loss"] for x in m["sched"]]
    m["con"] = [-x * I["concessions"] for x in m["sched"]]
    m["oi"] = [u * I["other_income_mo"] * 12 * ex[y] for y in Y]
    m["egi"] = [m["sched"][y] + m["vac"][y] + m["cl"][y] + m["con"][y] + m["oi"][y] for y in Y]

    def line(base):
        return [-u * base * ex[y] for y in Y]
    m["payroll"], m["repairs"], m["turnover"] = line(I["payroll"]), line(I["repairs"]), line(I["turnover_cost"])
    m["contract"], m["util"] = line(I["contract_svc"]), line(I["utilities"])
    ins_base = I["ins_roof"] if I["roof_toggle"] == 1 else I["ins_base"]
    m["ins"] = [-u * ins_base * (1 + I["ins_growth"]) ** y for y in Y]
    tax_y1 = P * I["cos_factor"] * I["millage"]
    m["tax"] = [-tax_y1 * (1 + I["nonhs_cap"]) ** y for y in Y]
    m["mgmt"] = [-m["egi"][y] * I["mgmt_fee_pct"] for y in Y]
    m["ga"], m["mktg"] = line(I["ga"]), line(I["marketing"])
    m["nav"] = line(I["nav"])  # per-unit non-ad valorem charges: not value-based
    opex_keys = ["payroll", "repairs", "turnover", "contract", "util", "ins", "tax", "nav", "mgmt", "ga", "mktg"]
    m["opex"] = [sum(m[k][y] for k in opex_keys) for y in Y]
    m["noi"] = [m["egi"][y] + m["opex"][y] for y in Y]
    m["noi_pretax"] = [m["noi"][y] - m["tax"][y] for y in Y]
    m["reserves"] = [-u * I["reserves"] * ex[y] for y in Y]
    m["capex"] = [-(I["roof_toggle"] * I["roof_cost"] if y == 0 else 0.0)
                  - ((I["recert_inspect"] + I["recert_remed"]) if (y + 1) == I["recert_year"] else 0.0) for y in Y]
    m["ucf"] = [m["noi"][y] + m["reserves"][y] + m["capex"][y] for y in Y]

    # Day-0 bridge and cap rates
    oi0 = u * I["other_income_mo"] * 12
    loss_pct = I["vacancy"] + I["credit_loss"] + I["concessions"]
    m["day0_gpr"] = sum(t["classic"] * t["F"] + t["renov"] * t["G"] for t in I["types"]) * 12
    m["day0_egi"] = m["day0_gpr"] * (1 - loss_pct) + oi0
    fixed0 = sum(m[k][0] for k in ["payroll", "repairs", "turnover", "contract", "util", "ins", "tax", "nav", "ga", "mktg"])
    m["day0_noi"] = m["day0_egi"] + fixed0 - m["day0_egi"] * I["mgmt_fee_pct"]
    s2_gpr = (sum(t["classic"] * t["F"] for t in I["types"])
              + sum(t["renov"] * (t["G"] + I["turnover_rate"] * I["burnoff_pct"] * (mk[t["name"]][0] - t["G"]))
                    for t in I["types"])) * 12
    s2_egi = s2_gpr * (1 - loss_pct) + oi0
    m["s2_noi"] = s2_egi + fixed0 - s2_egi * I["mgmt_fee_pct"]
    m["burnoff_effect"] = m["s2_noi"] - m["day0_noi"]
    m["capture_effect"] = m["noi"][0] - m["s2_noi"]
    m["cap_broker"] = I["broker_cap"]
    m["cap_inplace"] = m["day0_noi"] / P
    m["cap_y1fwd"] = m["noi"][0] / P
    m["exit_cap"] = m["cap_inplace"] + I["exit_spread_base"] if exit_cap is None else exit_cap
    m["exit_cap_sens"] = m["cap_inplace"] + I["exit_spread_sens"]
    m["eff_tax_exit"] = I["millage"] * I["cos_factor"]
    m["fwd_noi_pretax"] = m["noi_pretax"][HOLD]  # Year 6
    m["exit_price"] = m["fwd_noi_pretax"] / (m["exit_cap"] + m["eff_tax_exit"])
    m["cost_of_sale"] = -m["exit_price"] * I["cos_exit"]
    m["net_unlev"] = m["exit_price"] + m["cost_of_sale"]

    # Sources & uses (Scenario B base) and unlevered returns
    reno_per_unit = sum(I["reno_lines"]) * (1 + I["reno_contg_pct"])
    m["reno_total"] = reno_per_unit * I["classic_units"]
    m["closing"] = P * I["closing_pct"]
    m["uses_b"] = P + m["closing"] + m["reno_total"]
    m["unlev_cf"] = [-m["uses_b"]] + [m["ucf"][y] for y in range(HOLD)]
    m["unlev_cf"][HOLD] += m["net_unlev"]
    m["unlev_irr"] = _irr(m["unlev_cf"])
    m["unlev_em"] = sum(m["unlev_cf"][1:]) / m["uses_b"]

    # ---- Scenario B: new debt
    B = {}
    noi1 = m["noi"][0]
    B["loan_ltv"], B["loan_dy"], B["loan_dscr"] = P * I["ltv"], noi1 / I["min_dy"], noi1 / (I["min_dscr"] * rate)
    B["loan"] = min(B["loan_ltv"], B["loan_dy"], B["loan_dscr"])
    B["binding"] = ("LTV" if B["loan"] == B["loan_ltv"] else "Debt Yield" if B["loan"] == B["loan_dy"] else "DSCR")
    B["pmt"] = _pmt(rate, I["amort_years"], B["loan"])
    B["beg"], B["int"], B["prin"], B["ds"], B["end"] = [], [], [], [], []
    bal = B["loan"]
    for y in range(1, HOLD + 1):
        i_ = bal * rate
        p_ = 0.0 if y <= I["io_years"] else B["pmt"] - i_
        B["beg"].append(bal); B["int"].append(i_); B["prin"].append(p_); B["ds"].append(i_ + p_)
        bal -= p_
        B["end"].append(bal)
    B["equity"] = m["uses_b"] - B["loan"]
    B["lp_eq"], B["gp_eq"] = B["equity"] * (1 - I["gp_coinvest"]), B["equity"] * I["gp_coinvest"]
    B["payoff"] = -B["end"][-1]
    B["net_proceeds"] = m["exit_price"] + m["cost_of_sale"] + B["payoff"]
    B["lev_ops"] = [m["ucf"][y] - B["ds"][y] for y in range(HOLD)]
    B["lev_cf"] = [-B["equity"]] + list(B["lev_ops"])
    B["lev_cf"][HOLD] += B["net_proceeds"]
    B["irr"], B["em"] = _irr(B["lev_cf"]), sum(B["lev_cf"][1:]) / B["equity"]
    B["coc"] = [x / B["equity"] for x in B["lev_ops"]]

    # ---- Scenario A: assumed first + supplemental, refinance after maturity
    A = {}
    mat = I["a_mat"]
    for key, bal0, r_ in (("first", I["a_first_bal"], I["a_first_rate"]), ("supp", I["a_supp_bal"], I["a_supp_rate"])):
        beg, int_, prin, end = [], [], [], []
        bal = bal0
        pmt_ = _pmt(r_, I["amort_years"], bal0)
        for y in range(1, HOLD + 1):
            if y <= mat:
                i_ = bal * r_
                p_ = 0.0 if I["a_io"] == 1 else pmt_ - i_
                beg.append(bal); int_.append(i_); prin.append(p_)
                bal -= p_
                end.append(bal)
            else:
                beg.append(0.0); int_.append(0.0); prin.append(0.0); end.append(0.0)
        A[key] = dict(beg=beg, int=int_, prin=prin, end=end)
    A["payoff"] = A["first"]["end"][mat - 1] + A["supp"]["end"][mat - 1]
    A["refi_noi"] = m["noi"][mat]  # year mat+1
    A["refi_value"] = A["refi_noi"] / m["exit_cap"]  # appraisal at refi: in-place NOI / market cap (no reassessment)
    A["refi_ltv"] = (A["refi_value"] if refi_basis == "value" else P) * I["ltv"]  # "price" = pre-2026-09-29 rule
    A["refi_dy"] = A["refi_noi"] / I["min_dy"]
    A["refi_dscr"] = A["refi_noi"] / (I["min_dscr"] * rate)
    A["refi_loan"] = min(A["refi_ltv"], A["refi_dy"], A["refi_dscr"])
    A["refi_paydown"] = A["payoff"] - A["refi_loan"]
    A["refi_pmt"] = _pmt(rate, I["amort_years"], A["refi_loan"])
    beg, int_, prin, end = [], [], [], []
    bal = 0.0
    for y in range(1, HOLD + 1):
        if y > mat:
            if y == mat + 1:
                bal = A["refi_loan"]
            i_ = bal * rate
            p_ = 0.0 if (y - mat) <= I["io_years"] else A["refi_pmt"] - i_
            beg.append(bal); int_.append(i_); prin.append(p_)
            bal -= p_
            end.append(bal)
        else:
            beg.append(0.0); int_.append(0.0); prin.append(0.0); end.append(0.0)
    A["refi"] = dict(beg=beg, int=int_, prin=prin, end=end)
    A["tot_int"] = [A["first"]["int"][y] + A["supp"]["int"][y] + A["refi"]["int"][y] for y in range(HOLD)]
    A["tot_prin"] = [A["first"]["prin"][y] + A["supp"]["prin"][y] + A["refi"]["prin"][y] for y in range(HOLD)]
    A["tot_ds"] = [A["tot_int"][y] + A["tot_prin"][y] for y in range(HOLD)]
    A["tot_end"] = [A["first"]["end"][y] + A["supp"]["end"][y] + A["refi"]["end"][y] for y in range(HOLD)]
    A["fee"] = (I["a_first_bal"] + I["a_supp_bal"]) * I["a_fee_pct"]
    A["uses"] = m["uses_b"] + A["fee"]
    A["loan"] = I["a_first_bal"] + I["a_supp_bal"]
    A["equity"] = A["uses"] - A["loan"]
    A["lp_eq"], A["gp_eq"] = A["equity"] * (1 - I["gp_coinvest"]), A["equity"] * I["gp_coinvest"]
    A["payoff_exit"] = -A["tot_end"][-1]
    A["net_proceeds"] = m["exit_price"] + m["cost_of_sale"] + A["payoff_exit"]
    A["refi_short"] = [-(A["refi_paydown"] if (y + 1) == mat else 0.0) for y in range(HOLD)]
    A["lev_ops"] = [m["ucf"][y] - A["tot_ds"][y] + A["refi_short"][y] for y in range(HOLD)]
    A["lev_cf"] = [-A["equity"]] + list(A["lev_ops"])
    A["lev_cf"][HOLD] += A["net_proceeds"]
    A["irr"], A["em"] = _irr(A["lev_cf"]), sum(A["lev_cf"][1:]) / A["equity"]
    A["coc1"] = A["lev_ops"][0] / A["equity"]
    A["capital_call"] = max(0.0, -min(A["lev_ops"]))  # largest single-year shortfall funded by investors
    A["peak_equity"] = A["equity"] + A["capital_call"]

    # ---- Waterfalls (pari passu investor capital + GP promote), both scenarios
    W = waterfall(A["lev_cf"], A["equity"], I)
    B["W"] = waterfall(B["lev_cf"], B["equity"], I)

    m["B"], m["A"], m["W"] = B, A, W
    m["price"] = P
    return m


# ---------------------------------------------------------------------------
# Comparison against the recalculated workbook
# ---------------------------------------------------------------------------
class Comparator:
    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path, data_only=True)
        self.rows = []  # (sheet, line, max_diff, tol, n_points, ok, detail)

    def ws(self, name):
        return self.wb[name]

    def _record(self, sheet, line, mine, theirs, tol):
        diffs = []
        for i, (a, b) in enumerate(zip(mine, theirs)):
            if not isinstance(b, (int, float)):
                diffs.append((i, a, b, float("inf")))
            else:
                diffs.append((i, a, b, abs(a - b)))
        worst = max(diffs, key=lambda d: d[3])
        ok = worst[3] <= tol
        self.rows.append((sheet, line, worst[3], tol, len(diffs), ok, worst))

    def series(self, sheet, label, mine, cols, tol=TOL_USD, exact=False, after=None, col_label=1):
        ws = self.ws(sheet)
        start = _find_row(ws, after) + 1 if after else 1
        r = _find_row(ws, label, col=col_label, exact=exact, start=start)
        theirs = [ws.cell(row=r, column=c).value for c in cols]
        self._record(sheet, label.strip(), mine, theirs, tol)

    def scalar(self, sheet, label, mine, col=2, tol=TOL_USD, exact=False, after=None):
        self.series(sheet, label, [mine], [col], tol=tol, exact=exact, after=after)

    def text(self, sheet, label, mine, col=2):
        ws = self.ws(sheet)
        theirs = ws.cell(row=_find_row(ws, label), column=col).value
        self.rows.append((sheet, label, 0.0 if theirs == mine else float("inf"), 0, 1, theirs == mine,
                          (0, mine, theirs, 0)))


def compare(path, I, m):
    C = Comparator(path)
    OP, Y10 = "Operating Model", list(range(2, 2 + NYEARS))  # B..K = Years 1..10
    Y5 = list(range(3, 3 + HOLD))      # C..G = Years 1..5 on Returns / Returns (Assumed) / Waterfall
    Y05 = [2] + Y5                     # B..G = Years 0..5
    D5 = list(range(2, 2 + HOLD))      # B..F = Years 1..5 on Debt (Assumed)
    B, A, W = m["B"], m["A"], m["W"]

    # units consistency (the model's Units cell is a formula to the Unit Mix total)
    C.scalar("Assumptions", "Units (formula", I["units"], col=3, tol=0)

    # Operating model, every line, 10 years
    for t in I["types"]:
        n = t["name"]
        C.series(OP, f"{n} — Market Rent/mo", m["track_market"][n], Y10)
        C.series(OP, f"{n} — Classic-Unit Track Rent/mo", m["track_classic"][n], Y10)
        C.series(OP, f"{n} — Prior-Renovated Track Rent/mo", m["track_prior"][n], Y10)
    for lab, key, ex in [("Gross Potential Rent at Market", "market_gpr", False), ("Less: Loss-to-Lease", "ltl", False),
                         ("Gross Scheduled Rent", "sched", False), ("Less: Physical Vacancy", "vac", False),
                         ("Less: Credit Loss", "cl", False), ("Less: Concessions", "con", False),
                         ("Plus: Other Income", "oi", False), ("Effective Gross Income (EGI)", "egi", False),
                         ("Payroll", "payroll", True), ("Repairs & Maintenance", "repairs", True),
                         ("Turnover / Make-Ready", "turnover", True), ("Contract Services", "contract", True),
                         ("Utilities (owner-paid)", "util", True), ("Insurance", "ins", True),
                         ("Property Tax (reassessed basis", "tax", False), ("Non-Ad Valorem Assessments", "nav", False),
                         ("Management Fee (% of EGI)", "mgmt", True),
                         ("G&A / Admin", "ga", True), ("Marketing", "mktg", True),
                         ("Total Operating Expenses", "opex", True), ("NET OPERATING INCOME (NOI)", "noi", True),
                         ("memo: NOI Before Property Tax", "noi_pretax", False),
                         ("Less: Replacement Reserves", "reserves", True), ("Less: One-Time Capex", "capex", False),
                         ("UNLEVERED CASH FLOW", "ucf", False)]:
        C.series(OP, lab, m[key], Y10, exact=ex)

    # NOI bridge, cap rates, exit (Returns tab)
    R = "Returns"
    C.scalar(R, "Day-0 Gross Potential Rent", m["day0_gpr"])
    C.scalar(R, "Day-0 Effective Gross Income", m["day0_egi"])
    C.scalar(R, "DAY-0 IN-PLACE NOI", m["day0_noi"])
    C.scalar(R, "S2 NOI", m["s2_noi"])
    C.scalar(R, "+ Loss-to-lease burn-off effect", m["burnoff_effect"])
    C.scalar(R, "+ Renovation-premium partial-year capture effect", m["capture_effect"])
    C.scalar(R, "= YEAR-1 FORWARD NOI", m["noi"][0])
    C.scalar(R, "1. Broker-Stated Cap Rate", m["cap_broker"], tol=TOL_RATE)
    C.scalar(R, "2. In-Place Cap Rate", m["cap_inplace"], tol=TOL_RATE)
    C.scalar(R, "3. Year-1 FORWARD Cap Rate", m["cap_y1fwd"], tol=TOL_RATE)
    C.scalar(R, "Entry Cap Rate ANCHOR", m["cap_inplace"], tol=TOL_RATE)
    C.scalar(R, "Exit Cap Rate — base case", m["exit_cap"], tol=TOL_RATE)
    C.scalar(R, "Exit Cap Rate — sensitivity ceiling", m["exit_cap_sens"], tol=TOL_RATE)
    C.scalar(R, "Effective Tax Rate at Exit", m["eff_tax_exit"], tol=TOL_RATE)
    C.scalar(R, "Forward NOI Before Property Tax", m["fwd_noi_pretax"])
    C.scalar(R, "EXIT PRICE", m["exit_price"])
    C.scalar(R, "Less: Cost of Sale", m["cost_of_sale"])
    C.scalar(R, "Net Sale Proceeds, UNLEVERED", m["net_unlev"])
    C.series(R, "TOTAL UNLEVERED CASH FLOW TO EQUITY", m["unlev_cf"], Y05)
    C.scalar(R, "Unlevered IRR (debt-agnostic)", m["unlev_irr"], tol=TOL_RATE)
    C.scalar(R, "Unlevered Equity Multiple (debt-agnostic)", m["unlev_em"], tol=TOL_RATE)

    # Scenario B -- Capital, Debt, Returns
    C.scalar("Capital", "Purchase Price", m["price"])
    C.scalar("Capital", "Closing Costs", m["closing"])
    C.scalar("Capital", "Renovation Capex", m["reno_total"])
    C.scalar("Capital", "TOTAL USES", m["uses_b"])
    C.scalar("Capital", "Senior Loan", B["loan"])
    C.scalar("Capital", "Sponsor Equity (plug)", B["equity"])
    C.scalar("Capital", "LP Equity", B["lp_eq"], exact=True)
    C.scalar("Capital", "GP Equity (co-invest)", B["gp_eq"])
    C.scalar("Debt", "Loan Amount — LTV Constraint", B["loan_ltv"], col=3)
    C.scalar("Debt", "Loan Amount — Debt Yield Constraint", B["loan_dy"], col=3)
    C.scalar("Debt", "Loan Amount — DSCR Constraint", B["loan_dscr"], col=3)
    C.scalar("Debt", "SIZED LOAN AMOUNT", B["loan"], col=3)
    C.text("Debt", "Binding Constraint", B["binding"], col=3)
    C.scalar("Debt", "Annual P&I Payment (post-IO)", B["pmt"])
    dws = C.ws("Debt")
    yr_rows = {dws.cell(row=r, column=1).value: r for r in range(1, dws.max_row + 1)
               if isinstance(dws.cell(row=r, column=1).value, int)}
    for col, key, lab in [(2, "beg", "Beg. Balance"), (3, "int", "Interest"), (4, "prin", "Principal"),
                          (5, "ds", "Debt Service"), (6, "end", "End Balance")]:
        theirs = [dws.cell(row=yr_rows[y], column=col).value for y in range(1, HOLD + 1)]
        C._record("Debt", f"Scenario B schedule: {lab}", B[key], theirs, TOL_USD)
    C.series(R, "Less: Debt Service (Scenario B)", [-x for x in B["ds"]], Y5)
    C.series(R, "Levered CF from Operations (Scenario B)", B["lev_ops"], Y5)
    C.scalar(R, "Less: Loan Payoff (Scenario B", B["payoff"])
    C.scalar(R, "NET SALE PROCEEDS TO EQUITY — SCENARIO B", B["net_proceeds"])
    C.series(R, "TOTAL LEVERED CASH FLOW TO EQUITY — SCENARIO B", B["lev_cf"], Y05)
    C.scalar(R, "Levered IRR — SCENARIO B", B["irr"], tol=TOL_RATE)
    C.scalar(R, "Levered Equity Multiple — SCENARIO B", B["em"], tol=TOL_RATE)
    C.series(R, "CASH-ON-CASH BY YEAR", B["coc"], Y5, tol=TOL_RATE, after=None)  # label row itself is a header
    # (CoC values sit on the row after the header -- re-point)
    C.rows.pop()
    rws = C.ws(R)
    hr = _find_row(rws, "CASH-ON-CASH BY YEAR")
    C._record(R, "Cash-on-cash by year (Scenario B)", B["coc"], [rws.cell(row=hr + 1, column=c).value for c in Y5],
              TOL_RATE)

    # Scenario A -- Debt (Assumed)
    DA = "Debt (Assumed)"
    for sect, key in [("FIRST MORTGAGE", "first"), ("SUPPLEMENTAL LOAN", "supp"), ("POST-REFINANCE LOAN", "refi")]:
        C.series(DA, "Beginning Balance", A[key]["beg"], D5, after=sect)
        C.series(DA, "Interest", A[key]["int"], D5, exact=True, after=sect)
        C.series(DA, "Principal", A[key]["prin"], D5, after=sect)
        C.series(DA, "Ending Balance", A[key]["end"], D5, after=sect)
        C.rows[-4:] = [(s, f"{key}: {l}", d, t, n, ok, w) for (s, l, d, t, n, ok, w) in C.rows[-4:]]
    C.scalar(DA, "Payoff Balance", A["payoff"])
    C.scalar(DA, "NOI in the refinance year", A["refi_noi"])
    C.scalar(DA, "Value at Refinance", A["refi_value"])
    C.scalar(DA, "New Loan — LTV constraint", A["refi_ltv"])
    C.scalar(DA, "New Loan — Debt Yield constraint", A["refi_dy"])
    C.scalar(DA, "New Loan — DSCR constraint", A["refi_dscr"])
    C.scalar(DA, "NEW REFI LOAN AMOUNT", A["refi_loan"])
    C.scalar(DA, "Cash Required / (Released) at Refinance", A["refi_paydown"])
    C.scalar(DA, "Annual P&I Payment (post-refi", A["refi_pmt"])
    C.series(DA, "Total Interest", A["tot_int"], D5)
    C.series(DA, "Total Principal", A["tot_prin"], D5)
    C.series(DA, "TOTAL DEBT SERVICE", A["tot_ds"], D5)
    C.series(DA, "TOTAL ENDING BALANCE", A["tot_end"], D5)

    # Scenario A -- Returns (Assumed)
    RA = "Returns (Assumed)"
    C.scalar(RA, "Loan Assumption Fee", A["fee"])
    C.scalar(RA, "Total Uses (Scenario B uses + assumption fee)", A["uses"])
    C.scalar(RA, "Assumed Debt (First + Supplemental", A["loan"])
    C.scalar(RA, "Sponsor Equity", A["equity"], exact=True)
    C.scalar(RA, "LP Equity (90%)", A["lp_eq"])
    C.scalar(RA, "GP Equity (10%", A["gp_eq"])
    C.scalar(RA, "Exit Price (same tax-adjusted", m["exit_price"])
    C.scalar(RA, "Less: Cost of Sale", m["cost_of_sale"])
    C.scalar(RA, "Less: Loan Payoff (Debt (Assumed)", A["payoff_exit"])
    C.scalar(RA, "NET SALE PROCEEDS TO EQUITY — SCENARIO A", A["net_proceeds"])
    C.series(RA, "Unlevered CF (same as Scenario B", [m["ucf"][y] for y in range(HOLD)], Y5)
    C.series(RA, "Less: Debt Service (Debt (Assumed))", [-x for x in A["tot_ds"]], Y5)
    C.series(RA, "Less: Cash Required at Refinance", A["refi_short"], Y5)
    C.series(RA, "Levered CF from Operations", A["lev_ops"], Y5, exact=True)
    C.series(RA, "TOTAL LEVERED CASH FLOW TO EQUITY — SCENARIO A", A["lev_cf"], Y05)
    C.scalar(RA, "Levered IRR — SCENARIO A", A["irr"], tol=TOL_RATE)
    C.scalar(RA, "Levered Equity Multiple — SCENARIO A", A["em"], tol=TOL_RATE)
    C.scalar(RA, "Year-1 Cash-on-Cash — SCENARIO A", A["coc1"], tol=TOL_RATE)
    C.scalar(RA, "Capital Call Required", A["capital_call"])
    C.scalar(RA, "PEAK EQUITY — SCENARIO A", A["peak_equity"])

    # Waterfall (Scenario A), tier by tier -- engine uses a different (closed-form) method
    WF = "Waterfall"
    C.series(WF, "Available Cash for Distribution", W["avail"], Y5)
    C.scalar(WF, "Investor Capital, Year 0", A["equity"])
    C.series(WF, "Investor Pref Hurdle Balance, End of Year", W["pref_bal"], Y05)
    C.series(WF, "Tier 1 Distribution to Investors", W["t1"], Y5)
    C.series(WF, "Remaining Cash After Tier 1", W["rem1"], Y5)
    C.series(WF, "Investor Tier-2 Hurdle Balance, End of Year", W["h_bal"], Y05)
    C.series(WF, "Tier-2 Hurdle Balance After Tier-1 Cash Applied", W["h_after_t1"], Y5)
    C.series(WF, "Tier 2 Total Distribution", W["t2"], Y5)
    C.series(WF, "Tier 2 to investors", W["inv2"], Y5)
    C.series(WF, "Tier 2 GP promote", W["pr2"], Y5)
    C.series(WF, "Remaining Cash After Tier 2", W["rem2"], Y5)
    C.series(WF, "Tier 3 Total Distribution", W["t3"], Y5)
    C.series(WF, "Tier 3 to investors", W["inv3"], Y5)
    C.series(WF, "Tier 3 GP promote", W["pr3"], Y5)
    C.series(WF, "Total to Investors", [0.0] + W["inv_total"], Y05)
    C.series(WF, "LP Total Distribution", [-W["lp_eq"]] + W["lp"], Y05)
    C.series(WF, "GP Co-Invest Distribution", [-W["gp_eq"]] + W["gp_co"], Y05)
    C.series(WF, "GP Promote (Tier 2 + Tier 3)", [0.0] + W["promote"], Y05)
    C.series(WF, "GP Total Distribution", [-W["gp_eq"]] + W["gp"], Y05)
    C.scalar(WF, "LP IRR", W["lp_irr"], tol=TOL_RATE, exact=True)
    C.scalar(WF, "GP IRR (co-invest + promote)", W["gp_irr"], tol=TOL_RATE)
    C.scalar(WF, "LP Equity Multiple", W["lp_em"], tol=TOL_RATE, exact=True)
    C.scalar(WF, "GP Equity Multiple (co-invest + promote)", W["gp_em"], tol=TOL_RATE)
    C.scalar(WF, "Total GP Promote over the hold", W["promote_total"])

    # Summary headline cells (what gets quoted)
    SM = "Summary"
    C.scalar(SM, "Levered IRR (deal-level, Scenario A)", A["irr"], tol=TOL_RATE)
    C.scalar(SM, "Levered Equity Multiple (deal-level, Scenario A)", A["em"], tol=TOL_RATE)
    C.scalar(SM, "Unlevered IRR (debt-agnostic)", m["unlev_irr"], tol=TOL_RATE)
    C.scalar(SM, "LP IRR (Scenario A)", W["lp_irr"], tol=TOL_RATE)
    C.scalar(SM, "GP IRR (Scenario A)", W["gp_irr"], tol=TOL_RATE)
    C.scalar(SM, "2. In-Place Cap Rate", m["cap_inplace"], tol=TOL_RATE)
    C.scalar(SM, "Exit Cap Rate (base case", m["exit_cap"], tol=TOL_RATE)
    C.scalar(SM, "Sponsor Equity", A["equity"], exact=True)
    return C


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else "Parkview_Crossing_Acquisition_Model.xlsx"
    I = load_inputs(path)
    m = run_model(I)
    C = compare(path, I, m)
    fails = [r for r in C.rows if not r[5]]
    n_points = sum(r[4] for r in C.rows)
    print(f"verify_model.py -- {path}")
    print(f"Inputs read: {sum(1 for k in I if k not in ('types',))} scalar/list inputs + {len(I['types'])} unit types "
          f"(all asserted literal, none a formula). Units = {I['units']:.0f}.\n")
    print(f"{'Sheet':18s} {'Line':62s} {'pts':>4s} {'max |diff|':>14s}  status")
    for sheet, line, d, tol, n, ok, worst in C.rows:
        dstr = "inf" if d == float("inf") else (f"{d:.6f}" if tol < 1 else f"{d:,.4f}")
        print(f"{sheet[:18]:18s} {line[:62]:62s} {n:>4d} {dstr:>14s}  {'OK' if ok else 'DIFF'}")
    print(f"\n{len(C.rows)} lines, {n_points} individual figures compared "
          f"(tolerance: ${TOL_USD:.0f} on dollar lines, {TOL_RATE*1e4:.0f} bp on rates/IRRs/multiples).")
    if fails:
        print(f"RESULT: {len(fails)} LINE(S) DIFFER BEYOND TOLERANCE")
        for sheet, line, d, tol, n, ok, (i, mine, theirs, dd) in fails:
            print(f"  {sheet} / {line}: point {i}: mine={mine!r} workbook={theirs!r}")
        sys.exit(1)
    print("RESULT: every compared figure ties to the workbook within tolerance.")
    print(f"\nHeadline (recomputed independently): unlevered IRR {m['unlev_irr']:.2%} | "
          f"Scen A levered IRR {m['A']['irr']:.2%} / {m['A']['em']:.2f}x | "
          f"Scen B levered IRR {m['B']['irr']:.2%} / {m['B']['em']:.2f}x | "
          f"LP IRR {m['W']['lp_irr']:.2%} / {m['W']['lp_em']:.2f}x | GP IRR {m['W']['gp_irr']:.2%} | promote ${m['W']['promote_total']:,.0f}")


if __name__ == "__main__":
    main()
